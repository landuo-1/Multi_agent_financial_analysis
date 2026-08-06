"""
报告导出模块
支持 PDF、CSV、Excel 格式导出
"""

import os
import csv
import json
import time
from typing import Any, Dict, List, Optional
from src.utils.logger import get_logger

logger = get_logger(__name__)


def export_markdown_to_pdf(markdown_text: str, output_path: str, stock_code: str = "") -> str:
    """
    将 Markdown 报告转换为 PDF

    优先使用 weasyprint（高质量），降级到 fpdf2（轻量）
    """
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    # 方案1：weasyprint（最佳质量，支持中文）
    try:
        return _export_pdf_weasyprint(markdown_text, output_path)
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"weasyprint 导出失败: {e}")

    # 方案2：markdown + weasyprint
    try:
        return _export_pdf_markdown_html(markdown_text, output_path)
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"markdown+weasyprint 导出失败: {e}")

    # 方案3：fpdf2（轻量方案）
    try:
        return _export_pdf_fpdf(markdown_text, output_path)
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"fpdf2 导出失败: {e}")

    # 全部失败，保存为纯文本
    txt_path = output_path.replace(".pdf", ".txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(markdown_text)
    logger.warning(f"PDF 导出失败，已保存为纯文本: {txt_path}")
    return txt_path


def _export_pdf_weasyprint(markdown_text: str, output_path: str) -> str:
    """使用 weasyprint 导出 PDF（最佳质量）"""
    import markdown
    from weasyprint import HTML

    # Markdown → HTML
    html_content = markdown.markdown(
        markdown_text,
        extensions=["tables", "fenced_code", "codehilite"],
    )

    # 包装完整 HTML（含样式）
    full_html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            body {{ font-family: "Microsoft YaHei", "SimHei", sans-serif; margin: 40px; line-height: 1.8; }}
            h1 {{ color: #1a5276; border-bottom: 2px solid #1a5276; padding-bottom: 10px; }}
            h2 {{ color: #2874a6; border-bottom: 1px solid #ddd; padding-bottom: 5px; }}
            h3 {{ color: #3498db; }}
            table {{ border-collapse: collapse; width: 100%; margin: 15px 0; }}
            th, td {{ border: 1px solid #ddd; padding: 8px 12px; text-align: left; }}
            th {{ background-color: #f2f2f2; font-weight: bold; }}
            tr:nth-child(even) {{ background-color: #f9f9f9; }}
            code {{ background-color: #f4f4f4; padding: 2px 6px; border-radius: 3px; }}
            blockquote {{ border-left: 4px solid #3498db; margin: 15px 0; padding: 10px 20px; background: #f8f9fa; }}
            .disclaimer {{ background: #fff3cd; border: 1px solid #ffc107; padding: 15px; margin-top: 30px; border-radius: 5px; }}
            img {{ max-width: 100%; height: auto; }}
        </style>
    </head>
    <body>{html_content}</body>
    </html>
    """

    HTML(string=full_html).write_pdf(output_path)
    logger.info(f"PDF 导出成功: {output_path}")
    return output_path


def _export_pdf_markdown_html(markdown_text: str, output_path: str) -> str:
    """使用 markdown + weasyprint 导出"""
    import markdown
    from weasyprint import HTML

    html_content = markdown.markdown(markdown_text, extensions=["tables", "fenced_code"])
    HTML(string=html_content).write_pdf(output_path)
    return output_path


def _export_pdf_fpdf(markdown_text: str, output_path: str) -> str:
    """使用 fpdf2 导出（轻量方案）"""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # 添加中文字体
    font_path = _find_chinese_font()
    if font_path:
        pdf.add_font("Chinese", "", font_path)
        pdf.set_font("Chinese", size=10)
    else:
        pdf.set_font("Helvetica", size=10)

    page_width = pdf.w - 2 * pdf.l_margin  # 可用宽度

    # 逐行写入
    for line in markdown_text.split("\n"):
        stripped = line.strip()
        if not stripped:
            pdf.ln(3)
            continue

        if stripped.startswith("# "):
            pdf.set_font_size(16)
            pdf.multi_cell(page_width, 9, stripped[2:])
            pdf.ln(2)
            pdf.set_font_size(10)
        elif stripped.startswith("## "):
            pdf.set_font_size(13)
            pdf.multi_cell(page_width, 8, stripped[3:])
            pdf.ln(2)
            pdf.set_font_size(10)
        elif stripped.startswith("### "):
            pdf.set_font_size(11)
            pdf.multi_cell(page_width, 7, stripped[4:])
            pdf.ln(1)
            pdf.set_font_size(10)
        elif stripped.startswith("---"):
            pdf.ln(2)
            pdf.line(pdf.l_margin, pdf.y, pdf.w - pdf.r_margin, pdf.y)
            pdf.ln(2)
        elif stripped.startswith("|"):
            # 表格行（简化处理）
            pdf.multi_cell(page_width, 6, stripped)
        elif stripped.startswith(">"):
            # 引用块
            pdf.set_x(pdf.l_margin + 5)
            pdf.multi_cell(page_width - 5, 6, stripped[1:].strip())
        else:
            pdf.multi_cell(page_width, 6, stripped)

    pdf.output(output_path)
    logger.info(f"PDF 导出成功 (fpdf2): {output_path}")
    return output_path


def _find_chinese_font() -> Optional[str]:
    """查找系统中文字体"""
    import platform
    if platform.system() == "Windows":
        paths = [
            "C:/Windows/Fonts/msyh.ttc",    # 微软雅黑
            "C:/Windows/Fonts/simhei.ttf",   # 黑体
            "C:/Windows/Fonts/simsun.ttc",   # 宋体
        ]
    else:
        paths = [
            "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
            "/System/Library/Fonts/PingFang.ttc",
        ]
    for p in paths:
        if os.path.exists(p):
            return p
    return None


# ============================================================
# CSV / Excel 导出
# ============================================================

def export_financial_data_csv(agent_outputs: Dict[str, Any], output_path: str,
                              data_freshness: str = "",
                              data_period: str = "") -> str:
    """将分析数据导出为 CSV"""
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    import datetime
    export_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

    rows = []
    for agent_id, output in agent_outputs.items():
        agent_name = output.get("agent_name", agent_id)
        fact_sheet = output.get("fact_sheet", {})

        for key, value in fact_sheet.items():
            if key in ("risk_signals", "风险信号", "risk_alerts", "风险预警"):
                for signal in (value or []):
                    rows.append({
                        "导出时间": export_time,
                        "数据年份": data_period,
                        "数据截止": data_freshness,
                        "分析模块": agent_name,
                        "指标": "风险信号",
                        "数值": str(signal),
                        "状态": "warning",
                    })
            else:
                if isinstance(value, list):
                    value = "、".join(str(x) for x in value if x)
                rows.append({
                    "导出时间": export_time,
                    "数据年份": data_period,
                    "数据截止": data_freshness,
                    "分析模块": agent_name,
                    "指标": key,
                    "数值": value,
                    "状态": output.get("status", "unknown"),
                })

    if not rows:
        logger.warning("无数据可导出")
        return ""

    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["导出时间", "数据年份", "数据截止", "分析模块", "指标", "数值", "状态"])
        writer.writeheader()
        writer.writerows(rows)

    logger.info(f"CSV 导出成功: {output_path} ({len(rows)} 行)")
    return output_path


def export_financial_data_excel(agent_outputs: Dict[str, Any], output_path: str,
                                debate_consensus: Dict[str, Any] = None,
                                data_freshness: str = "",
                                data_period: str = "") -> str:
    """将分析数据导为 Excel（含多个 Sheet）"""
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        logger.warning("openpyxl 未安装，降级为 CSV")
        return export_financial_data_csv(agent_outputs, output_path.replace(".xlsx", ".csv"), data_freshness)

    import datetime
    export_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

    wb = openpyxl.Workbook()

    # Sheet1: 各 Agent 详细数据
    ws1 = wb.active
    ws1.title = "分析数据"

    # 表头样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2874A6", end_color="2874A6", fill_type="solid")

    headers = ["导出时间", "数据年份", "数据截止", "分析模块", "指标", "数值", "状态"]
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    row = 2
    for agent_id, output in agent_outputs.items():
        agent_name = output.get("agent_name", agent_id)
        fact_sheet = output.get("fact_sheet", {})
        status = output.get("status", "unknown")

        for key, value in fact_sheet.items():
            if key in ("risk_signals", "风险信号", "risk_alerts", "风险预警"):
                for signal in (value or []):
                    ws1.cell(row=row, column=1, value=export_time)
                    ws1.cell(row=row, column=2, value=data_period)
                    ws1.cell(row=row, column=3, value=data_freshness)
                    ws1.cell(row=row, column=4, value=agent_name)
                    ws1.cell(row=row, column=5, value="风险信号")
                    ws1.cell(row=row, column=6, value=str(signal))
                    ws1.cell(row=row, column=7, value="warning")
                    row += 1
            else:
                # 处理 list 值
                if isinstance(value, list):
                    value = "、".join(str(x) for x in value if x)
                ws1.cell(row=row, column=1, value=export_time)
                ws1.cell(row=row, column=2, value=data_period)
                ws1.cell(row=row, column=3, value=data_freshness)
                ws1.cell(row=row, column=4, value=agent_name)
                ws1.cell(row=row, column=5, value=key)
                ws1.cell(row=row, column=6, value=str(value))
                ws1.cell(row=row, column=7, value=status)
                row += 1

    # 自动列宽
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws1.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    # Sheet2: 辩论结论
    if debate_consensus:
        ws2 = wb.create_sheet("辩论结论")
        ws2.cell(row=1, column=1, value="项目").font = header_font
        ws2.cell(row=1, column=1).fill = header_fill
        ws2.cell(row=1, column=2, value="内容").font = header_font
        ws2.cell(row=1, column=2).fill = header_fill

        items = [
            ("多方置信度", debate_consensus.get("final_bull_confidence", "")),
            ("空方置信度", debate_consensus.get("final_bear_confidence", "")),
            ("胜方", debate_consensus.get("winner", "")),
            ("辩论轮次", debate_consensus.get("debate_rounds", "")),
            ("风险提示", debate_consensus.get("risk_warning", "")),
        ]
        for i, (k, v) in enumerate(items, 2):
            ws2.cell(row=i, column=1, value=k)
            ws2.cell(row=i, column=2, value=str(v))

        # 共识点
        row = len(items) + 3
        ws2.cell(row=row, column=1, value="共识点").font = Font(bold=True)
        for point in debate_consensus.get("consensus", []):
            row += 1
            ws2.cell(row=row, column=1, value=f"- {point}")

    wb.save(output_path)
    logger.info(f"Excel 导出成功: {output_path}")
    return output_path


def export_all(stock_code: str, stock_name: str,
               final_report: str, agent_outputs: Dict[str, Any],
               debate_consensus: Dict[str, Any] = None,
               output_dir: str = "./reports") -> Dict[str, str]:
    """
    一键导出所有格式
    返回 {format: filepath}
    """
    os.makedirs(output_dir, exist_ok=True)
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    base_name = f"{stock_code}_{timestamp}"

    results = {}

    # 1. Markdown
    md_path = os.path.join(output_dir, f"{base_name}.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(final_report)
    results["markdown"] = md_path

    # 2. PDF
    pdf_path = os.path.join(output_dir, f"{base_name}.pdf")
    try:
        export_markdown_to_pdf(final_report, pdf_path, stock_code)
        results["pdf"] = pdf_path
    except Exception as e:
        logger.error(f"PDF 导出失败: {e}")

    # 3. Excel
    xlsx_path = os.path.join(output_dir, f"{base_name}.xlsx")
    try:
        export_financial_data_excel(agent_outputs, xlsx_path, debate_consensus)
        results["excel"] = xlsx_path
    except Exception as e:
        logger.error(f"Excel 导出失败: {e}")

    # 4. CSV
    csv_path = os.path.join(output_dir, f"{base_name}.csv")
    try:
        export_financial_data_csv(agent_outputs, csv_path)
        results["csv"] = csv_path
    except Exception as e:
        logger.error(f"CSV 导出失败: {e}")

    logger.info(f"导出完成: {list(results.keys())}")
    return results
